import sys
import os
import re
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, '.')

from src.utils.paddle_ocr import run_pdf_text_extraction
from src.utils.tag_classifier import classify_paddle_results, canonicalize_tag
from src.agents.compiler import CompilerAgent

wb = openpyxl.load_workbook(r'pid_stuff/PID_Extraction_RevB_Corrected_26-KA-901_902 (2).xlsx', data_only=True)

def parse_gt(dwg_key):
    gt = {'EQUIPMENT': [], 'PSV': [], 'INSTRUMENTS': [], 'VALVES': [], 'LINES': []}
    for r in list(wb['Equipment List'].iter_rows(values_only=True))[3:16]:
        if not r or r[1] is None:
            continue
        tag = str(r[1]).strip()
        if dwg_key == '901' and ('901' in tag or tag in ('26-CK-911', '26-CX-9122', '26-HA-911')):
            gt['EQUIPMENT'].append(tag)
        elif dwg_key == '902' and ('902' in tag or tag in ('26-CK-921', '26-CX-9222')):
            gt['EQUIPMENT'].append(tag)

    for r in list(wb['Safety Relief Valve List'].iter_rows(values_only=True))[3:7]:
        if not r or r[1] is None:
            continue
        tag = str(r[1]).strip()
        if dwg_key == '901' and '9066' in tag:
            gt['PSV'].append(tag)
        elif dwg_key == '902' and '9027' in tag:
            gt['PSV'].append(tag)

    curr = None
    for r in list(wb['Instrument List'].iter_rows(values_only=True)):
        if not r:
            continue
        c0 = str(r[0]).strip() if r[0] is not None else ''
        c1 = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ''
        if '26-KA-901 INSTRUMENTS' in c0:
            curr = '901'
            continue
        if '26-KA-902 INSTRUMENTS' in c0:
            curr = '902'
            continue
        if curr == dwg_key and c1 and not c1.startswith('#') and 'Tag' not in c1:
            gt['INSTRUMENTS'].append(c1)

    curr = None
    for r in list(wb['Manual Valve List'].iter_rows(values_only=True)):
        if not r:
            continue
        c0 = str(r[0]).strip() if r[0] is not None else ''
        c1 = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ''
        if '26-KA-901 MANUAL VALVES' in c0:
            curr = '901'
            continue
        if '26-KA-902 MANUAL VALVES' in c0:
            curr = '902'
            continue
        if curr == dwg_key and c1 and not c1.startswith('#') and 'Tag' not in c1:
            gt['VALVES'].append(c1)

    curr = None
    for r in list(wb['Line List'].iter_rows(values_only=True)):
        if not r:
            continue
        c0 = str(r[0]).strip() if r[0] is not None else ''
        c1 = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ''
        if '26-KA-901 LINES' in c0:
            curr = '901'
            continue
        if '26-KA-902 LINES' in c0:
            curr = '902'
            continue
        if curr == dwg_key and c1 and not c1.startswith('#') and 'Line Number' not in c1:
            gt['LINES'].append(c1)

    return gt

def clean_tag(t):
    t = re.sub(r'^\d{2,3}-', '', t.upper().strip())
    for ch in [' ', '"', "'", '-']:
        t = t.replace(ch, '')
    return t

def evaluate_drawing(dwg_name, pdf_path, dwg_key):
    gt = parse_gt(dwg_key)
    print(f"\n=======================================================")
    print(f"EVALUATING {dwg_name} ({pdf_path})")
    print(f"=======================================================")
    ocr_items = run_pdf_text_extraction(pdf_path)
    classified = classify_paddle_results(ocr_items, 'PID')

    ca = CompilerAgent()
    state = {
        'extracted_entities': {
            'text_elements': classified,
            'symbols': [],
            'geometry': {},
            'relations': [],
        },
        'drawing_type': 'PID',
        'cv_results': {},
        'revision_history': []
    }
    res = ca.run(state)
    graph = res['engineering_graph']

    compiled_dict = {
        'EQUIPMENT': graph.equipment,
        'LINES': graph.lines,
        'INSTRUMENTS': graph.instruments,
        'VALVES': graph.valves,
        'PSV': graph.safety_relief_valves
    }

    results = {}
    tot_tp, tot_fp, tot_fn, tot_gt, tot_comp = 0, 0, 0, 0, 0

    for cat, gt_list in gt.items():
        comp_list = [item.tag for item in compiled_dict[cat]]
        gt_map = {clean_tag(t): t for t in gt_list}
        comp_map = [clean_tag(t) for t in comp_list]

        tp_gt = set()
        matched_comp = set()

        for c_idx, c_tag in enumerate(comp_map):
            if c_tag in gt_map:
                tp_gt.add(c_tag)
                matched_comp.add(c_idx)
            else:
                for g_tag in gt_map:
                    if (len(c_tag) >= 4 and c_tag in g_tag) or (len(g_tag) >= 4 and g_tag in c_tag):
                        tp_gt.add(g_tag)
                        matched_comp.add(c_idx)
                        break

        tp = len(tp_gt)
        fp = len(comp_list) - len(matched_comp)
        fn = len(gt_list) - tp
        rec = tp / len(gt_list) if gt_list else 0.0
        prec = len(matched_comp) / len(comp_list) if comp_list else 0.0
        f1 = 2 * prec * rec / (prec + rec) if (prec + rec) else 0.0

        missed = [t for t in gt_list if clean_tag(t) not in tp_gt]
        extra = [comp_list[i] for i in range(len(comp_list)) if i not in matched_comp]

        results[cat] = {
            'gt': len(gt_list),
            'comp': len(comp_list),
            'tp': tp,
            'fp': fp,
            'fn': fn,
            'rec': round(rec, 4),
            'prec': round(prec, 4),
            'f1': round(f1, 4),
            'missed': missed,
            'extra': extra
        }

        tot_tp += tp
        tot_comp += len(comp_list)
        tot_gt += len(gt_list)
        tot_fp += fp
        tot_fn += fn

        print(f"{cat:12s} | GT: {len(gt_list):2d} | Comp: {len(comp_list):2d} | TP: {tp:2d} | FP: {fp:2d} | FN: {fn:2d} | Rec: {rec*100:6.2f}% | Prec: {prec*100:6.2f}% | F1: {f1:.4f}")
        if missed:
            print(f"   Missed GT: {missed}")
        if extra:
            print(f"   Extra Comp ({len(extra)}): {extra[:8]}{'...' if len(extra)>8 else ''}")

    overall_rec = tot_tp / tot_gt if tot_gt else 0.0
    overall_prec = (tot_comp - tot_fp) / tot_comp if tot_comp else 0.0
    overall_f1 = 2 * overall_prec * overall_rec / (overall_prec + overall_rec) if (overall_prec + overall_rec) else 0.0

    print("-" * 65)
    print(f"OVERALL {dwg_name}: GT: {tot_gt} | Comp: {tot_comp} | TP: {tot_tp} | FP: {tot_fp} | FN: {tot_fn} | Recall: {overall_rec*100:.2f}% | Precision: {overall_prec*100:.2f}% | F1: {overall_f1:.4f}")
    print(f"Generic components count: {len(graph.generic_components)}")
    if graph.generic_components:
        print(f"   Generic tags: {[g.tag for g in graph.generic_components[:10]]}")
    print(f"Relationships count: {len(graph.relationships)}")
    return {
        'overall': {
            'gt': tot_gt, 'comp': tot_comp, 'tp': tot_tp, 'fp': tot_fp, 'fn': tot_fn,
            'recall': overall_rec, 'precision': overall_prec, 'f1': overall_f1,
            'generic_count': len(graph.generic_components),
            'rel_count': len(graph.relationships)
        },
        'categories': results
    }

if __name__ == '__main__':
    res_a = evaluate_drawing("DRAWING A (26-KA-901)", r'pid_stuff/Lift Gas compressor-PID.pdf', '901')
    res_b = evaluate_drawing("DRAWING B (26-KA-902)", r'pid_stuff/Export Gas Compressor-PID.pdf', '902')

    with open('PRECISION_DUAL_EVAL_BEFORE.json', 'w', encoding='utf-8') as f:
        json.dump({'drawing_a': res_a, 'drawing_b': res_b}, f, indent=2)
    print("\nSaved dual evaluation results to PRECISION_DUAL_EVAL_BEFORE.json")
